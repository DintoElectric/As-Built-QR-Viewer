#!/usr/bin/env python3
"""
Import per-J-box circuit lists from the circuits spreadsheet into
public/data/drawings.json, attaching a `circuits` array to each J-box callout.

Matches each spreadsheet row to a drawing callout on (panel, tag#), where the
drawing tag# is the parenthetical (stored as the jbox `ckts` field, joined by
commas). Rows on the spreadsheet's "Disregarded" tab are ignored.

    python3 scripts/import_circuits.py path/to/2nd_floor_jbox_circuits.xlsx

Requires: openpyxl  (pip install openpyxl)
"""
import json, os, sys, re
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DRAWINGS = os.path.join(ROOT, 'public/data/drawings.json')

def norm(s):
    return re.sub(r'\s+', '', str(s)) if s is not None else ''

def main(xlsx):
    wb = load_workbook(xlsx, data_only=True)
    ws = wb['2nd floor J-box circuits'] if '2nd floor J-box circuits' in wb.sheetnames else wb.worksheets[0]
    # header row = the one containing 'J-box'
    hr = next((r for r in range(1, 12) if ws.cell(r, 1).value == 'J-box'), 4)
    lut = {}                         # (panel, tag#) -> [circuits]
    for r in range(hr + 1, ws.max_row + 1):
        jb, panel = ws.cell(r, 1).value, ws.cell(r, 2).value
        if jb and panel and 'callout' in str(jb):   # group header
            continue
        if not panel:
            continue
        tag = norm(ws.cell(r, 3).value)
        circ = ws.cell(r, 4).value
        if not circ:
            continue
        circuits = [c.strip() for c in str(circ).replace(';', ',').split(',') if c.strip()]
        lut[(str(panel).strip(), tag)] = circuits

    data = json.load(open(DRAWINGS))
    matched = 0
    for sheet in data['sheets']:
        for b in sheet['jboxes']:
            key = (b['panel'], norm(','.join(b['ckts'])))
            if key in lut:
                b['circuits'] = lut[key]
                matched += 1
            elif 'circuits' in b:
                del b['circuits']
    json.dump(data, open(DRAWINGS, 'w'))
    print(f'spreadsheet rows with circuits: {len(lut)} | drawing callouts matched: {matched}')

if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, '2nd_floor_jbox_circuits.xlsx'))
